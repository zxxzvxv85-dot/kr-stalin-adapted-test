"""Create the illustrated education decision and full USSR poster catalogue.

Reads the supplied archive without extracting its paths. Does not modify game files.
"""
from __future__ import annotations
import argparse
import hashlib
import io
import json
import math
from pathlib import Path
import zipfile

from PIL import Image, ImageOps
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from education_reform_content import ACTIVITIES, CONTROLS, POSTER_CATEGORIES
from education_poster_expansion import STATE_REWARDS,MILITARY_GATE,FINAL_STRUGGLE_GATE,VICTORY_GATE

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT/'tools/design/education_system_20261024/poster_catalog.json'
MODIFIERS = {
 'research_speed_factor':'科研速度', 'industrial_capacity_factory':'工厂产出',
 'stability_factor':'稳定度', 'consumer_goods_factor':'生活消费品工厂系数',
 'political_power_factor':'政治点获取', 'political_power_cost':'每日政治点消耗',
 'army_defence_factor':'陆军师防御','army_attack_factor':'陆军师攻击','army_morale_factor':'陆军组织度恢复',
 'army_org_factor':'陆军组织度','planning_speed':'计划速度','war_support_factor':'战争支持度',
 'production_factory_efficiency_gain_factor':'生产效率增长','production_factory_start_efficiency_factor':'基础生产效率',
 'production_speed_buildings_factor':'建造速度','production_speed_infrastructure_factor':'基础设施建造速度',
 'enemy_operative_detection_chance_factor':'敌方特工被发现几率','conscription_factor':'适役人口系数',
 'training_time_army_factor':'陆军训练时间','local_org_regain':'地区组织度恢复',
 'state_production_speed_buildings_factor':'地区建造速度','local_building_slots_factor':'地区最大工厂数',
 'local_resources_factor':'地区资源获取',
 'monthly_population':'每月人口','equipment_conversion_speed':'装备改装速度',
 'political_advisor_cost_factor':'政治顾问花费','stability_weekly':'每周稳定度',
 'resistance_growth_on_our_occupied_states':'我方占领地区抵抗增长','production_speed_rail_way_factor':'铁路建设速度',
 'production_speed_supply_node_factor':'补给中心建设速度','production_factory_max_efficiency_factor':'生产效率上限',
 'production_speed_bunker_factor':'陆上要塞建设速度','experience_loss_factor':'有经验的士兵损失',
 'drift_defence_factor':'意识形态变化抵制力度','mobilization_speed':'动员速度',
 'conversion_cost_civ_to_mil_factor':'民用工厂转军用工厂花费','conversion_cost_mil_to_civ_factor':'军用工厂转民用工厂花费',
 'production_lack_of_resource_penalty_factor':'缺乏战略资源惩罚','supply_consumption_factor':'补给消耗',
 'industrial_capacity_dockyard':'海军船坞产出','max_planning_factor':'计划加成上限',
 'experience_gain_army_factor':'陆军经验增长','command_power_gain_mult':'指挥点数增长',
}

def modifier_text(values):
    return '；'.join(f'{MODIFIERS[k]} {v:+g}' if k=='political_power_cost' else f'{MODIFIERS[k]} {v*100:+g}%' for k,v in values.items()) or '无'

def configure(doc):
    section=doc.sections[0];section.page_width=Inches(8.2677);section.page_height=Inches(11.6929)
    section.top_margin=section.bottom_margin=Inches(.60)
    section.left_margin=section.right_margin=Inches(.65)
    for node in list(doc.styles.element.iter(qn('w:pBdr'))):
        node.getparent().remove(node)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
        st=doc.styles[name];st.font.name='Microsoft YaHei';st.font.color.rgb=RGBColor(0,0,0)
        st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Microsoft YaHei')
    doc.styles['Normal'].font.size=Pt(10)
    doc.styles['Normal'].paragraph_format.space_after=Pt(5)
    doc.styles['Normal'].paragraph_format.line_spacing=1.12
    doc.styles['Title'].font.size=Pt(24)
    doc.styles['Subtitle'].font.italic=False
    doc.styles['Heading 1'].font.size=Pt(17)
    doc.styles['Heading 2'].font.size=Pt(12)
    doc.core_properties.title='教育改革宣传决议与苏联海报图录'
    doc.core_properties.author=''
    footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=footer.add_run();r.font.size=Pt(9)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE')
    cached=OxmlElement('w:r');text=OxmlElement('w:t');text.text='1';cached.append(text);field.append(cached);footer._p.append(field)

def table(doc, rows, widths, header=False):
    t=doc.add_table(rows=rows,cols=len(widths));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    pr=t._tbl.tblPr;borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        n=OxmlElement('w:'+edge);n.set(qn('w:val'),'single');n.set(qn('w:sz'),'5');n.set(qn('w:color'),'D9D9D9');borders.append(n)
    pr.append(borders)
    for row in t.rows:
        no_split=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no_split)
        for cell,width in zip(row.cells,widths):
            cell.width=Inches(width);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            margins=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),'90');e.set(qn('w:type'),'dxa');margins.append(e)
            cell._tc.get_or_add_tcPr().append(margins)
    if header:
        repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
        for c in t.rows[0].cells:
            fill=OxmlElement('w:shd');fill.set(qn('w:fill'),'E9EEF2');c._tc.get_or_add_tcPr().append(fill)
    return t

def para(cell,text='',bold=False,size=10):
    p=cell.add_paragraph();r=p.add_run(text);r.bold=bold;r.font.size=Pt(size);return p

def image(cell,data,title,maxw,maxh):
    im=Image.open(io.BytesIO(data)).convert('RGB');im=ImageOps.contain(im,(900,900))
    f=io.BytesIO();im.save(f,format='JPEG',quality=88);f.seek(0)
    scale=min(maxw/im.width,maxh/im.height)
    p=cell.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    shape=p.add_run().add_picture(f,width=Inches(im.width*scale),height=Inches(im.height*scale))
    shape._inline.docPr.set('descr',title)

def build(archive, output):
    manifest=json.loads(DATA.read_text(encoding='utf-8'))
    assert hashlib.sha256(archive.read_bytes()).hexdigest()==manifest['sha256'],'Archive differs from recorded source'
    posters=manifest['posters'];by_code={p['code']:p for p in posters}
    assert len(posters)==len(by_code)==271
    adopted=len({a['poster'] for a in ACTIVITIES});unused=len(posters)-adopted
    doc=Document();configure(doc)
    doc.add_paragraph('教育改革宣传决议与苏联海报图录','Title')
    doc.add_paragraph('测试版现有宣传效果与完整素材目录','Subtitle')
    doc.add_paragraph(f'本册列出教育改革系统现有的{len(ACTIVITIES)}项海报宣传与2项热情调节，并按六类收录03_USSR素材包的全部271张海报。已采用海报标明对应决议；其余标明未采用及效果未配置。')
    doc.add_paragraph('素材中文标题用于检索，以画面和原文件名拟定，不作为经过考证的历史正式译名。游戏宣传口号另列，保留模组背景下的改编名称。')
    doc.add_heading('六类目录',1)
    t=table(doc,7,[1.65,1.25,1.1,2.9],True)
    for c,s in zip(t.rows[0].cells,['分类','原文件夹','素材数量','已实装宣传数量']):c.text=s
    for row,(cat,(name,folders)) in zip(t.rows[1:],POSTER_CATEGORIES.items()):
        vals=[name,' '.join(f'{x:02}' for x in folders),str(sum(p['category']==cat for p in posters)),str(sum(a['category']==cat for a in ACTIVITIES))]
        for c,s in zip(row.cells,vals):c.text=s
    doc.add_heading('阅读说明',1)
    doc.add_paragraph('决议页中的教育增量为基础总增量，实际提升受热情和教育路线影响；每次宣传加成后的总提升仍不超过基础教育4、高等教育2个百分点。热情变化按每30天列示。正常完成宣传次数加1，取消不增加次数。每类教学活动最多并行1项，总栏位为2项，完成对文化革命后为3项。')
    doc.add_paragraph('六类按钮筛选决议与当前在办海报，未开始的活动不张贴。基础教育、高等教育、职业技术、文化和国防动员各最多并行1项，热情调节决议在所有分类中保留。素材图录按原文件夹分类，游戏决议分类按主题指定。')
    doc.add_paragraph('宣传的国家收益与代价在活动期间同时生效，结束或取消时移除。原有永久奖励按既定数额保留，新增加的临时效果与代价不转为永久修正。教育进度与正常完成次数保留；指定地区、农业奖励及第聂伯河攻势按各条目列示。')
    with zipfile.ZipFile(archive) as z:
        for start in range(0,len(ACTIVITIES),3):
            doc.add_page_break();doc.add_heading(f'已实装宣传决议 {start//3+1}',1)
            for a in ACTIVITIES[start:start+3]:
                item=by_code[a['poster']];t=table(doc,1,[2.35,4.55]);left,right=t.rows[0].cells
                image(left,z.read(item['file']),item['title_zh'],2.10,2.25)
                p=left.add_paragraph(f"{item['code']}  {item['title_zh']}");p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                right.paragraphs[0].add_run(a['name']).bold=True
                para(right,'分类  '+POSTER_CATEGORIES[a['category']][0])
                para(right,f"花费 {a['cost']}政治点    持续 {a['days']}天    冷却90天")
                gains=[]
                if a['basic']:gains.append(f"基础教育 +{a['basic']:g}")
                if a['higher']:gains.append(f"高等教育 +{a['higher']:g}")
                para(right,'教育基础增量  '+('；'.join(gains) or '无'))
                para(right,f"热情变化  每30天 {a['heat']:+g}")
                para(right,'期间效果  '+modifier_text(a['modifiers']))
                if a['permanent_modifiers']:para(right,'完成后原有永久效果  '+modifier_text(a['permanent_modifiers']))
                access={'basic':'完成共产主义教育理念','higher':'完成对文化革命且基础教育至少35','technical':'完成专业技能培训','culture':'完成对文化革命','military':'完成共产主义教育理念'}[a['kind']]
                if a.get('gate'):access+='；'+{MILITARY_GATE:'完成军改第四阶段',FINAL_STRUGGLE_GATE:'完成最后的斗争',VICTORY_GATE:'完成胜利日'}[a['gate']]
                if a.get('wartime'):access+='；正在交战（停战取消）'
                if a.get('state_reward'):
                    r=a['state_reward'];name,values=STATE_REWARDS[r['preset']]
                    state_name={195:'彼得格勒／列宁格勒',219:'莫斯科',202:'基辅'}[r['state']]
                    para(right,f'完成地区效果  {state_name}：{name}，{r["days"]}天；'+modifier_text(values),size=9)
                    access+='；拥有并控制'+state_name
                if a.get('agriculture_reward'):
                    text={'wheat':'小麦库存+4（上限12）','rye':'黑麦库存+3（上限12）','technical':'本季作物产量+5%','repair':'本季农机损耗率−5个百分点（不降低已有较强支援）'}[a['agriculture_reward']]
                    para(right,'国家农业启用且未关闭时，完成后：'+text,size=9)
                if a.get('timed_reward'):
                    r=a['timed_reward']
                    para(right,f'完成后额外效果  对乌克兰核心领土攻击+{r["attack"]*100:g}%，持续{r["days"]}天',size=9)
                para(right,'解锁条件  '+access,size=9)
                doc.add_paragraph().paragraph_format.space_after=Pt(3)
        doc.add_page_break();doc.add_heading('热情调节与使用规则',1)
        for a in CONTROLS:
            doc.add_heading(a['name'].strip('！'),2)
            condition='热情低于35' if a['id']=='listen' else '热情至少70'
            doc.add_paragraph(f"条件 {condition}。花费20政治点，持续30天，冷却60天。期间每日政治点消耗减少{-a['modifiers']['political_power_cost']:g}；结束时热情{a['heat']:+g}，按原有奖励永久减少每日政治点消耗{-a['permanent_modifiers']['political_power_cost']:g}。")
        doc.add_paragraph('热情调节不占海报栏位，两项不能同时进行，不计完成宣传次数，也没有单独海报。')
        doc.add_heading('教育系统初值',2)
        doc.add_paragraph('基础教育46，高等教育7，民众热情32。热情每日自然下降0.10；两个教育水平不自然衰减。')
        doc.add_heading('热情与经济动员',2)
        doc.add_paragraph('工厂产出、基础生产效率及生产效率增长的正向收益随热情增强，科研等其他收益不统一放大。基础生产效率最高10%，效率增长最高12.5%；热情提供的工厂产出最高5%，另加已累计的永久贡献。')
        doc.add_paragraph('热情超过75后，每周损失稳定度，损失随热情连续加重：80热情时每周降低0.04个百分点，90时降低0.36，100时降低1。教育收益与其他过热惩罚仍各自生效。')
        doc.add_heading('国策追加效果',2)
        doc.add_paragraph('集中宣传供给保留原有收益，同一海报正常完成后的重复宣传花费、临时正负修正、教育及热情增量变为1.5倍；永久奖励不变，教育仍受4/2上限约束。鼓励业余创作保留原有收益，并将海报花费减半；先获得180天群众犹豫不决（每日额外热情−0.15），再获得180天群众热情澎湃（每日额外+0.05），之后结束。向左增加1个科研槽；时代召唤使海报花费再降低25%。费用优惠乘算，两项无海报热情调节不参与。')
        doc.add_heading('国策解锁要求',2)
        for line in ['第一组  完成宣传3次，基础教育至少30。','第二组  完成宣传6次，基础教育至少45，高等教育至少15。','第三组  完成宣传9次，基础教育至少60，高等教育至少25。']:
            doc.add_paragraph(line)
        doc.add_heading('素材采用状态',2)
        doc.add_paragraph(f'已实装{adopted}张素材、{len(ACTIVITIES)}项宣传；未采用{unused}张，均未配置决议效果。未采用素材只收录于后续图录，不加入GUI或可执行决议。')
        for category,(name,folders) in POSTER_CATEGORIES.items():
            entries=[p for p in posters if p['category']==category]
            for start in range(0,len(entries),6):
                doc.add_page_break();doc.add_heading(f'{category} {name} 海报图录 {start//6+1}',1)
                doc.add_paragraph('原文件夹 '+ ' '.join(f'{n:02}' for n in folders)+f'    本类共{len(entries)}张',style='Caption')
                selected=entries[start:start+6];t=table(doc,math.ceil(len(selected)/2),[3.45,3.45])
                for cell,item in zip((c for row in t.rows for c in row.cells),selected):
                    image(cell,z.read(item['file']),item['title_zh'],3.05,1.95)
                    p=para(cell,f"{item['code']}  {item['title_zh']}",True,10);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                    status='已实装  效果见前文同编号宣传决议' if item['activity'] else '未采用  效果未配置'
                    p=para(cell,status,False,9);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                    p=para(cell,f"原目录 {item['folder']:02}",False,8);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    output.parent.mkdir(parents=True,exist_ok=True);doc.save(output)
    print(f'Created {output} ({output.stat().st_size:,} bytes); 271 catalogue entries and {len(ACTIVITIES)} illustrated decisions.')

if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--archive',type=Path,default=Path('C:/Users/Administrator/Downloads/03_USSR.zip'))
    parser.add_argument('--output',type=Path,default=ROOT/'output/education_catalog_20261024/教育改革宣传决议与苏联海报图录.docx')
    args=parser.parse_args();build(args.archive,args.output)
